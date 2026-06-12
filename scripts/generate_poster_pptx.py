"""
生成 PPTX 格式研究海报 → report/poster.pptx
============================================
用法: conda run -n zf-li23 python scripts/generate_poster_pptx.py
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import os
os.environ['DISPLAY'] = ''  # 防止 GUI 报错

FIGS_DIR = Path(__file__).resolve().parent.parent / "report" / "figures"

# ── 颜色 ─────────────────────────────────────────────────────────────
DARK_BLUE = RGBColor(0x2C, 0x3E, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF8, 0xF9, 0xFA)
MID_GRAY = RGBColor(0xBD, 0xC3, 0xC7)
DARK_TEXT = RGBColor(0x2C, 0x3E, 0x50)
ACCENT_BLUE = RGBColor(0x34, 0x98, 0xDB)
ACCENT_GREEN = RGBColor(0x27, 0xAE, 0x60)

prs = Presentation()
# A1 竖版: 594 x 841 mm → 转换成 EMU (1mm = 36000 EMU)
SLIDE_W = Cm(59.4)
SLIDE_H = Cm(84.1)
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

# ── 辅助函数 ─────────────────────────────────────────────────────────

def add_rect(left, top, width, height, fill_color=LIGHT_GRAY, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.5)
    else:
        shape.line.fill.background()
    return shape

def add_textbox(left, top, width, height, text, fontsize=12, bold=False,
                color=DARK_TEXT, alignment=PP_ALIGN.LEFT, word_wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(fontsize)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return txBox

def add_image(path, left, top, width, height=None):
    if not path.exists():
        print(f"  ⚠ Image not found: {path}")
        return
    if height:
        slide.shapes.add_picture(str(path), left, top, width, height)
    else:
        slide.shapes.add_picture(str(path), left, top, width)


# ── 标题栏 ───────────────────────────────────────────────────────────

title_bar = add_rect(Cm(1), Cm(1), Cm(57.4), Cm(4.5), fill_color=DARK_BLUE)
add_textbox(Cm(2), Cm(1.3), Cm(55), Cm(2.5),
            'SubCellSpace: A Unified Database for Subcellular Spatial Transcriptomics',
            fontsize=28, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

add_textbox(Cm(2), Cm(3.5), Cm(55), Cm(1.5),
            'F. Li  ·  School of Life Sciences, Tsinghua University  ·  2026',
            fontsize=12, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

# ── 左栏: 背景 + 管道 ───────────────────────────────────────────────

col_x = Cm(1.5)
col_w = Cm(27)
y_start = Cm(6.5)

# Background box
add_rect(col_x, y_start, col_w, Cm(5.5), fill_color=WHITE)
add_textbox(col_x + Cm(1), y_start + Cm(0.3), col_w - Cm(2), Cm(0.8),
            'Background', fontsize=16, bold=True, color=DARK_BLUE)
add_textbox(col_x + Cm(1), y_start + Cm(1.3), col_w - Cm(2), Cm(4.0),
            'Spatial transcriptomics enables gene expression profiling within\n'
            'native tissue context. SubCellSpace builds a comprehensive,\n'
            'scverse-native database from three major platforms:\n'
            '  • NanoString CosMx — High-plex in situ RNA detection\n'
            '  • 10x Genomics Xenium — Subcellular + multi-modal imaging\n'
            '  • Vizgen MERFISH — Multiplexed FISH with ultra-high plexity',
            fontsize=9, color=DARK_TEXT)

# Pipeline figure
add_rect(col_x, y_start + Cm(6.0), col_w, Cm(8.5), fill_color=WHITE)
add_textbox(col_x + Cm(1), y_start + Cm(6.3), col_w - Cm(2), Cm(0.8),
            'Data Processing Pipeline', fontsize=16, bold=True, color=DARK_BLUE)
add_image(FIGS_DIR / 'fig5_pipeline.png', col_x + Cm(0.5), y_start + Cm(7.5),
          col_w - Cm(1), Cm(6.5))

# ── 中间 & 右栏: 数据库概览 ─────────────────────────────────────────

mid_x = Cm(30)
mid_w = Cm(28)
y_start2 = Cm(6.5)

# Overview figure
add_rect(mid_x, y_start2, mid_w, Cm(13), fill_color=WHITE)
add_textbox(mid_x + Cm(1), y_start2 + Cm(0.3), mid_w - Cm(2), Cm(0.8),
            'Database Scale & Enrichment', fontsize=16, bold=True, color=DARK_BLUE)
add_image(FIGS_DIR / 'fig1_overview.png', mid_x + Cm(0.3), y_start2 + Cm(1.5),
          mid_w - Cm(0.6), Cm(11))

# Stats summary
add_rect(mid_x, y_start2 + Cm(13.5), mid_w, Cm(3.5), fill_color=RGBColor(0xE8, 0xF8, 0xF0))
add_textbox(mid_x + Cm(1), y_start2 + Cm(14.0), mid_w - Cm(2), Cm(2.5),
            'Database Summary\n'
            'Total datasets: 2,131  |  Ready: 1,019\n'
            'CosMx: 476  |  Xenium: 1,525  |  MERFISH: 130\n'
            '21 data columns per entry: GEO metadata + LLM annotation\n'
            '+ file integrity (touch) + streaming statistics',
            fontsize=10, color=DARK_TEXT)

# Optional files
add_rect(mid_x, y_start2 + Cm(17.5), mid_w, Cm(8.5), fill_color=WHITE)
add_textbox(mid_x + Cm(1), y_start2 + Cm(17.8), mid_w - Cm(2), Cm(0.8),
            'Companion File Availability', fontsize=16, bold=True, color=DARK_BLUE)
add_image(FIGS_DIR / 'fig2_optional_files.png', mid_x + Cm(0.3), y_start2 + Cm(19.0),
          mid_w - Cm(0.6), Cm(6.5))

# Biological coverage
add_rect(col_x, y_start + Cm(15.0), col_w, Cm(10.5), fill_color=WHITE)
add_textbox(col_x + Cm(1), y_start + Cm(15.3), col_w - Cm(2), Cm(0.8),
            'Biological Coverage', fontsize=16, bold=True, color=DARK_BLUE)
add_image(FIGS_DIR / 'fig3_tissue_disease.png', col_x + Cm(0.3), y_start + Cm(16.5),
          col_w - Cm(0.6), Cm(8.5))

# CosMx stats
add_rect(col_x, y_start + Cm(26.0), col_w, Cm(10.5), fill_color=WHITE)
add_textbox(col_x + Cm(1), y_start + Cm(26.3), col_w - Cm(2), Cm(0.8),
            'Dataset Statistics (CosMx)', fontsize=16, bold=True, color=DARK_BLUE)
add_image(FIGS_DIR / 'fig4_stats_summary.png', col_x + Cm(0.3), y_start + Cm(27.5),
          col_w - Cm(0.6), Cm(8.5))

# ── 底部技术栈 ───────────────────────────────────────────────────────

bottom_bar = add_rect(Cm(1), Cm(80.5), Cm(57.4), Cm(2.5), fill_color=DARK_BLUE)
add_textbox(Cm(2), Cm(81.0), Cm(55), Cm(1.5),
            'Tech Stack: Python · scanpy · squidpy · spatialdata · React · TypeScript · Slurm\n'
            'Data: NCBI GEO · DeepSeek LLM · 10x Genomics · NanoString/Bruker · Vizgen',
            fontsize=10, color=MID_GRAY, alignment=PP_ALIGN.CENTER)


# ── 保存 ─────────────────────────────────────────────────────────────

output_path = Path(__file__).resolve().parent.parent / "report" / "poster.pptx"
prs.save(str(output_path))
print(f"  ✅ PPTX saved: {output_path}")
print(f"     Size: {output_path.stat().st_size / 1024:.0f} KB")
